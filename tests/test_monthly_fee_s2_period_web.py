"""
月費修正 S2（Web）— 收費期間（起算月份／最後收費月份）。

Web 透過 web.utils.db._build_connection 直接連線（不走桌面 _create_tables/_migrate），
所以測試分兩種資料庫：
- migrated：先由桌面 Database() 升級（模擬「桌面先更新、再同步到 Drive」）
- legacy  ：尚未升級的舊資料庫 → Web 必須維持舊行為、不報錯、不寫入
"""
import io
import os
import sqlite3
os.environ.setdefault("ADMIN_USERNAME", "")

import openpyxl
import pytest
from fastapi.testclient import TestClient

def _ym_iter(start, end):
    y, m = divmod(start, 100)
    while y * 100 + m <= end:
        yield y, m
        m += 1
        if m > 12:
            y, m = y + 1, 1


def _make_legacy(path):
    """完整的「S2 之前」資料庫：先用目前 Database() 建立全部資料表，
    再移除 S2 欄位，得到與舊版完全相同的結構（含 transactions 等其他表）。"""
    import database as database_module
    database_module.DB_PATH = path
    database_module.Database().close()
    c = sqlite3.connect(str(path))
    c.execute("ALTER TABLE rent_tenants DROP COLUMN start_ym")
    c.execute("ALTER TABLE rent_tenants DROP COLUMN end_ym")
    for n, g, a, ca in [("A老", "三年級", 1, "2026-06-10T10:00:00"),
                        ("B新", "一年級", 1, "2026-07-02T10:00:00"),
                        ("C停", "二年級", 0, "2026-06-10T10:00:00")]:
        c.execute("INSERT INTO rent_tenants(name,monthly_rent,grade,active,created_at) "
                  "VALUES(?,3000,?,?,?)", (n, g, a, ca))
    for tid, s, e in [(1, 202510, 202608), (2, 202607, 202608), (3, 202510, 202606)]:
        for y, m in _ym_iter(s, e):
            c.execute("INSERT INTO rent_payments(tenant_id,year,month,status,paid_date) "
                      "VALUES(?,?,?,'paid','2026-07-01')", (tid, y, m))
    c.commit()
    c.close()


def _client_for(path):
    from web.app import app
    from web.utils.db import get_db, ReadOnlyDB, _build_connection
    conn_db = _build_connection(str(path))       # Web 實際使用的連線方式

    def override():
        db = ReadOnlyDB.__new__(ReadOnlyDB)
        db._db = conn_db
        yield db

    app.dependency_overrides[get_db] = override
    return app, conn_db


@pytest.fixture
def migrated(tmp_path):
    import database as database_module
    path = tmp_path / "data.db"
    _make_legacy(path)
    database_module.DB_PATH = path
    database_module.Database().close()            # 桌面升級
    app, conn_db = _client_for(path)
    with TestClient(app) as c:
        yield c, conn_db
    app.dependency_overrides.clear()


@pytest.fixture
def legacy(tmp_path):
    path = tmp_path / "data.db"
    _make_legacy(path)
    app, conn_db = _client_for(path)
    with TestClient(app) as c:
        yield c, conn_db, path
    app.dependency_overrides.clear()


def _order(html, names):
    html = html[html.find("繳費明細"):]      # S15：只看繳費明細表（表格上方的逾期明細也會出現姓名）
    pos = {n: html.find(n) for n in names}
    return [n for n in sorted(pos, key=pos.get) if pos[n] >= 0]


def test_w1_rent_page_period(migrated):
    c, _ = migrated
    h6 = c.get("/rent?year=2026&month=6").text
    assert _order(h6, ["A老", "B新", "C停"]) == ["A老", "C停"]
    h7 = c.get("/rent?year=2026&month=7").text
    assert _order(h7, ["A老", "B新", "C停"]) == ["A老", "B新"]
    api = c.get("/api/rent/payments?year=2026&month=6").json()
    assert [(p["name"], p["status"]) for p in api] == [("A老", "paid"), ("C停", "paid")]


def test_w2_restricted_rent_page_period(migrated):
    c, _ = migrated
    r = c.get("/restricted/rent?year=2026&month=6")
    if r.status_code == 404:
        pytest.skip("此版本沒有 /restricted（完整版 v6）")
    assert _order(r.text, ["A老", "B新", "C停"]) == ["A老", "C停"]


def test_w3_legacy_db_keeps_old_behavior(legacy):
    c, _, _ = legacy
    for url in ("/rent?year=2026&month=6", "/tenants", "/",
                "/export/rent.xlsx?start_year=2026&start_month=5&end_year=2026&end_month=9"):
        assert c.get(url).status_code == 200, url
    names = [p["name"] for p in c.get("/api/rent/payments?year=2026&month=6").json()]
    assert sorted(names) == ["A老", "B新"]          # 舊行為：在職者全部列出，停用者不列


def test_w4_rent_excel_grey_outside_period(migrated):
    c, _ = migrated
    r = c.get("/export/rent.xlsx?start_year=2026&start_month=5&end_year=2026&end_month=9")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["月費繳交明細表"]
    rows = {ws.cell(i, 3).value: i for i in range(4, ws.max_row + 1)}
    assert set(rows) == {"A老", "B新"}
    fill = lambda n, j: ws.cell(rows[n], 4 + j).fill.fgColor.rgb[-6:]
    assert fill("B新", 0) == "CCCCCC" and fill("B新", 1) == "CCCCCC"   # 05、06 起算前
    assert ws.cell(rows["B新"], 6).value == "V"                        # 07
    assert fill("A老", 4) == "FFFFFF"                                   # 09 收費中、無紀錄
    ws3 = wb["停用人員紀錄"]
    r3 = [ws3.cell(i, 3).value for i in range(4, ws3.max_row + 1)]
    assert r3 == ["C停"]
    assert ws3.cell(4, 5).value == "V"                                   # 06
    assert ws3.cell(4, 6).fill.fgColor.rgb[-6:] == "CCCCCC"             # 07 停用後


def test_w5_web_never_writes_or_migrates(legacy):
    c, conn_db, path = legacy
    before = list(sqlite3.connect(str(path)).iterdump())
    for url in ("/rent?year=2026&month=6", "/tenants", "/", "/api/rent/tenants",
                "/api/rent/payments?year=2026&month=6",
                "/export/rent.xlsx?start_year=2026&start_month=5&end_year=2026&end_month=9"):
        assert c.get(url).status_code == 200
    after = list(sqlite3.connect(str(path)).iterdump())
    assert before == after
    cols = [r[1] for r in sqlite3.connect(str(path)).execute("PRAGMA table_info(rent_tenants)")]
    assert "start_ym" not in cols                                        # Web 沒有 ALTER


def test_w6_dashboard_source_uses_period(migrated):
    """Dashboard 本月月費小卡的合計直接來自 get_payments_for_month()（與 /rent 同一來源）。"""
    c, conn_db = migrated
    ps = conn_db.get_payments_for_month(2026, 6)
    assert sum(p["monthly_rent"] for p in ps if p["status"] != "exempt") == 6000
    assert c.get("/").status_code == 200
