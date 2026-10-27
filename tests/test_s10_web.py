"""
S10 Web：期初餘額／交接轉出
  收入不含期初、支出不含交接轉出；結餘＝期初＋收入－支出－轉出（數字與桌面相同）
  記帳列表標示「期初」「轉出」；統計表「收支差」
  仍然唯讀：讀到尚未升級的舊資料庫也不寫入、不升級
"""
import re
import sqlite3

import pytest


@pytest.fixture
def special(desktop_db):
    d = desktop_db                       # 已有 2026-07-01 收入 5000
    d.add_transaction("income", 10000, "期初餘額", "前期轉入", "2026-06-01")
    d.add_transaction("expense", 1200, "交通", "", "2026-07-05")
    d.add_transaction("expense", 8000, "交接轉出", "交給下一任", "2026-07-31")
    return d


def _ok(client, url):
    r = client.get(url)
    if r.status_code == 404:
        pytest.skip(f"{url} 不在這個版本")
    assert r.status_code == 200, url
    return r


def _money(text):
    return int(re.sub(r"[^\d\-]", "", text))


def test_summary_api_has_opening_and_balance(client, special):
    r = _ok(client, "/api/transactions/summary").json()
    assert r["total_income"] == 5000 and r["total_expense"] == 1200
    assert r["total_opening"] == 10000 and r["total_transfer_out"] == 8000
    assert r["balance"] == 10000 + 5000 - 1200 - 8000


def test_summary_matches_desktop(client, special):
    r = _ok(client, "/api/transactions/summary").json()
    assert {k: r[k] for k in special.get_summary()} == special.get_summary()


def test_dashboard_shows_opening_note_and_balance(client, special):
    h = _ok(client, "/").text
    assert "另有期初" in h and "另有交接轉出" in h
    m = re.search(r"累計結餘.*?stat-value[^>]*>\s*([^<]+)<", h, re.S)
    assert m and _money(m.group(1)) == 5800


def test_transactions_badges_and_totals(client, special):
    h = _ok(client, "/transactions?date_from=2026-01-01&date_to=2026-12-31").text
    assert 'badge bg-info text-dark">期初<' in h
    assert 'badge bg-warning text-dark">轉出<' in h
    m = re.search(r"區間收入.*?stat-value[^>]*>\s*([^<]+)<", h, re.S)
    assert m and _money(m.group(1)) == 5000
    m = re.search(r"區間結餘.*?stat-value[^>]*>\s*([^<]+)<", h, re.S)
    assert m and _money(m.group(1)) == 5800


def test_statistics_header_renamed_and_income_excludes_opening(client, special):
    h = _ok(client, "/statistics?year=2026").text
    assert "收支差" in h and "另有期初" in h
    m = re.search(r" 收入<.*?stat-value[^>]*>\s*([^<]+)<", h, re.S)
    assert m and _money(m.group(1)) == 5000


def test_search_summary_mentions_opening(client, special):
    h = _ok(client, "/search?q=前期轉入").text
    m = re.search(r"期初 <span class=\"fw-semibold\">([^<]+)<", h)    # 摘要列另外列出期初
    assert m and _money(m.group(1)) == 10000


def test_excel_export_year_sheet(client, special):
    import io
    import openpyxl
    r = _ok(client, "/export/transactions.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["年度統計"] if "年度統計" in wb.sheetnames else None
    if ws is None:
        pytest.skip("這個版本沒有年度統計工作表")
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    hdr = next(r for r in rows if r and r[0] == "年度")
    assert hdr[3].startswith("收支差")
    y = next(r for r in rows if r and str(r[0]) == "2026")
    assert y[1:4] == [5000, 1200, 3800]


def test_old_database_not_upgraded_by_web(client, desktop_db, tmp_path):
    """Web 讀到桌面還沒升級的舊資料庫：不建立分類、不改「前期轉入」、不寫入。"""
    import web.utils.db as wdb
    path = desktop_db.db_path
    c = sqlite3.connect(str(path))
    c.execute("DELETE FROM categories WHERE name IN ('期初餘額','交接轉出')")
    c.execute("INSERT INTO transactions(type,amount,category,note,date,created_at,voucher_no) "
              "VALUES('income',777,'其他收入','前期轉入','2026-06-01','x','11506001')")
    c.commit()
    before = list(c.iterdump())
    c.close()
    conn = wdb._build_connection(str(path))
    try:
        rows = conn.get_transactions() if hasattr(conn, "get_transactions") else None
    finally:
        try:
            conn.conn.close()
        except Exception:
            pass
    c = sqlite3.connect(str(path))
    assert list(c.iterdump()) == before
    c.close()
    if rows is not None:
        from database import summarize_rows
        s = summarize_rows(rows)
        assert s["total_opening"] == 0 and s["total_income"] == 5777   # 還沒升級：仍算收入
