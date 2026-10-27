"""
月費修正 S3（Web）— 標準月費／個人優惠／已繳金額固定。

Web 透過 web.utils.db._build_connection 直接連線（不走桌面遷移）：
- migrated：桌面 Database() 先升級，並新增「2026-12 起 3,500」、A 預繳 2026-12
- s2_only ：S3 尚未升級（沒有 rent_fee_* 表、沒有 amount 欄）→ Web 退回 monthly_rent
"""
import io
import os
import sqlite3
os.environ.setdefault("ADMIN_USERNAME", "")

import openpyxl
import pytest
from fastapi.testclient import TestClient


def _ym_iter(start, end):
    y, m = divmod(start, 100)
    while y * 100 + m <= end:
        yield y, m
        m += 1
        if m > 12:
            y, m = y + 1, 1


def _make_db(path, upgrade_s3=True):
    """A 3000、B 2500（→ 優惠 −500）、D 3000；A 2026-07～08 已繳，B 2026-07 已繳。"""
    import database as database_module
    database_module.DB_PATH = path
    d = database_module.Database()
    for n, g, rent in [("A", "三年級", 3000), ("B", "一年級", 2500), ("D", "一年級", 3000)]:
        d.conn.execute("INSERT INTO rent_tenants(name,monthly_rent,grade,active,created_at,start_ym) "
                       "VALUES(?,?,?,1,'2026-07-01T00:00:00',202607)", (n, rent, g))
    d.conn.commit()
    for tid, s, e in [(1, 202607, 202608), (2, 202607, 202607)]:
        for y, m in _ym_iter(s, e):
            d.set_payment_status(tid, y, m, "paid", paid_date="2026-07-01")
    if upgrade_s3:
        # 模擬「舊資料升級」：以遷移規則建立標準月費與優惠，再新增調價
        d.conn.execute("DELETE FROM rent_fee_schedule")
        d.conn.execute("DELETE FROM rent_fee_discount")
        d.conn.commit()
        d.add_fee_schedule(202607, 3000, "初始")
        d.set_discount(2, "amount", 500, 202607)
        d.set_payment_status(1, 2026, 12, "paid", paid_date="2026-10-01")   # A 預繳（3000）
        d.add_fee_schedule(202612, 3500, "調價")
    else:
        d.conn.execute("DROP TABLE rent_fee_schedule")
        d.conn.execute("DROP TABLE rent_fee_discount")
        d.conn.execute("ALTER TABLE rent_payments DROP COLUMN amount")
        d.conn.commit()
    d.close()


def _client(path):
    from web.app import app
    from web.utils.db import get_db, ReadOnlyDB, _build_connection
    conn_db = _build_connection(str(path))

    def override():
        db = ReadOnlyDB.__new__(ReadOnlyDB)
        db._db = conn_db
        yield db

    app.dependency_overrides[get_db] = override
    return app


@pytest.fixture
def migrated(tmp_path):
    path = tmp_path / "data.db"
    _make_db(path, upgrade_s3=True)
    app = _client(path)
    with TestClient(app) as c:
        yield c, path
    app.dependency_overrides.clear()


@pytest.fixture
def s2_only(tmp_path):
    path = tmp_path / "data.db"
    _make_db(path, upgrade_s3=False)
    app = _client(path)
    with TestClient(app) as c:
        yield c, path
    app.dependency_overrides.clear()


def _amounts(c, y, m):
    return {p["name"]: p["monthly_rent"]
            for p in c.get(f"/api/rent/payments?year={y}&month={m}").json()}


def test_w1_month_amounts_and_fixed_prepaid(migrated):
    c, _ = migrated
    assert _amounts(c, 2026, 11) == {"A": 3000, "B": 2500, "D": 3000}
    assert _amounts(c, 2026, 12) == {"A": 3000, "B": 3000, "D": 3500}   # A 預繳固定


def test_w2_rent_page_totals(migrated):
    c, _ = migrated
    h = c.get("/rent?year=2026&month=12").text
    assert "NT$9,500" in h            # 應收 3000 + 3000 + 3500
    assert "NT$3,000" in h            # 已收（A）


def test_w3_restricted_rent_page(migrated):
    c, _ = migrated
    r = c.get("/restricted/rent?year=2026&month=12")
    if r.status_code == 404:
        pytest.skip("此版本沒有 /restricted（完整版 v6）")
    assert "NT$9,500" in r.text


def test_w4_tenants_page_current_fee(migrated):
    c, _ = migrated
    t = {x["name"]: x for x in c.get("/api/rent/tenants").json()}
    assert t["B"]["discount_text"] == "−500"
    assert t["B"]["monthly_rent"] == t["A"]["monthly_rent"] - 500
    assert c.get("/tenants").status_code == 200


def test_w5_excel_amount_sheet(migrated):
    c, _ = migrated
    r = c.get("/export/rent.xlsx?start_year=2026&start_month=11&end_year=2026&end_month=12")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["月費金額明細"]
    rows = {ws.cell(i, 3).value: i for i in range(4, ws.max_row + 1)}
    assert ws.cell(rows["A"], 5).value == 3000          # 2026-12 預繳固定


def test_w6_s2_only_database_fallback(s2_only):
    c, path = s2_only
    assert _amounts(c, 2026, 7) == {"A": 3000, "B": 2500, "D": 3000}
    for url in ("/rent?year=2026&month=7", "/tenants", "/", "/api/rent/tenants",
                "/export/rent.xlsx?start_year=2026&start_month=7&end_year=2026&end_month=8"):
        assert c.get(url).status_code == 200, url


def test_w7_web_never_writes_or_migrates(s2_only):
    c, path = s2_only
    before = list(sqlite3.connect(str(path)).iterdump())
    for url in ("/rent?year=2026&month=7", "/tenants", "/", "/api/rent/tenants",
                "/export/rent.xlsx?start_year=2026&start_month=7&end_year=2026&end_month=8"):
        c.get(url)
    assert list(sqlite3.connect(str(path)).iterdump()) == before
    names = {r[0] for r in sqlite3.connect(str(path)).execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    assert "rent_fee_schedule" not in names
