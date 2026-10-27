"""
月費修正 S1（Web）— 月費人員年級排序 六→一→其他，同年級依姓名。

Web 透過 ReadOnlyDB 轉呼叫與桌面相同的 database.py，因此驗證的是
「實際 HTTP 回應內容」中的順序，而不是只測 SQL。
測試資料刻意打亂（一年級 id 最小、六年級 id 最大，含未填年級）。
"""
import io
import os
os.environ.setdefault("ADMIN_USERNAME", "")

import openpyxl
import pytest
from fastapi.testclient import TestClient

_SEED = [
    ("乙一", "一年級", 1), ("甲一", "一年級", 1), ("無年級者", "", 1),
    ("丙三", "三年級", 1), ("甲二", "二年級", 1), ("甲五", "五年級", 1),
    ("甲四", "四年級", 1), ("停用二", "二年級", 0), ("乙六", "六年級", 1),
    ("甲六", "六年級", 1),
]
_EXPECTED_ACTIVE_NAMES = ["乙六", "甲六", "甲五", "甲四", "丙三", "甲二", "乙一", "甲一", "無年級者"]


@pytest.fixture
def sort_client(tmp_path):
    import database as database_module
    database_module.DB_PATH = tmp_path / "s1_web.db"
    d = database_module.Database()
    for name, grade, active in _SEED:
        d.conn.execute(
            "INSERT INTO rent_tenants (name, monthly_rent, grade, active, created_at) "
            "VALUES (?, 3000, ?, ?, '2026-01-01T00:00:00')", (name, grade, active))
    d.conn.commit()
    d.set_payment_status(9, 2026, 7, "paid", paid_date="2026-07-01")  # 乙六
    d.set_payment_status(8, 2026, 7, "paid", paid_date="2026-07-01")  # 停用二

    from web.app import app
    from web.utils.db import get_db, ReadOnlyDB

    def override():
        db = ReadOnlyDB.__new__(ReadOnlyDB)
        db._db = d
        yield db

    app.dependency_overrides[get_db] = override
    with TestClient(app) as c:
        yield c, d
    app.dependency_overrides.clear()


def _positions(html, names):
    pos = [html.find(n) for n in names]
    assert all(p >= 0 for p in pos), dict(zip(names, pos))
    return pos


def test_w1_rent_page_html_order(sort_client):
    c, _ = sort_client
    r = c.get("/rent?year=2026&month=7")
    assert r.status_code == 200
    # S15（Q14 A）：未完成在上、已完成在下，兩區各自照年級；7 月只有乙六已繳 → 排到最後。
    # 只看繳費明細表（S15 起表格上方的逾期明細也會出現姓名）。
    html = r.text[r.text.find("繳費明細"):]
    expected = [n for n in _EXPECTED_ACTIVE_NAMES if n != "乙六"] + ["乙六"]
    pos = _positions(html, expected)
    assert pos == sorted(pos)


def test_w3_tenants_page_group_order(sort_client):
    c, _ = sort_client
    r = c.get("/tenants")
    assert r.status_code == 200
    pos = _positions(r.text, ["六年級", "五年級", "四年級", "三年級", "二年級", "一年級", "未分級"])
    assert pos == sorted(pos)


def test_w4_api_order(sort_client):
    c, _ = sort_client
    t = c.get("/api/rent/tenants").json()
    assert [x["name"] for x in t] == _EXPECTED_ACTIVE_NAMES
    p = c.get("/api/rent/payments?year=2026&month=7").json()
    assert [x["name"] for x in p] == _EXPECTED_ACTIVE_NAMES
    p6 = c.get("/api/rent/payments", params={"year": 2026, "month": 7, "grade": "六年級"}).json()
    assert [x["name"] for x in p6] == ["乙六", "甲六"]


def test_w5_rent_excel_row_order_rule(sort_client):
    """Web 月費 Excel 自行排序（在職 Sheet1/2 六→一，停用在 Sheet3），S1 不改其規則。"""
    c, _ = sort_client
    r = c.get("/export/rent.xlsx?start_year=2026&start_month=7&end_year=2026&end_month=7")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    for ws in wb.worksheets[:2]:
        names = [ws.cell(row, 3).value for row in range(4, ws.max_row + 1)]
        names = [n for n in names if n]
        assert names == _EXPECTED_ACTIVE_NAMES, ws.title


def test_w6_no_db_write(sort_client):
    c, d = sort_client
    dump = lambda: {t: d.conn.execute(f"SELECT * FROM {t} ORDER BY id").fetchall()
                    for t in ("rent_tenants", "rent_payments", "transactions")}
    before = [list(map(tuple, v)) for v in dump().values()]
    for url in ("/rent?year=2026&month=7", "/tenants", "/api/rent/tenants",
                "/api/rent/payments?year=2026&month=7",
                "/export/rent.xlsx?start_year=2026&start_month=7&end_year=2026&end_month=7"):
        assert c.get(url).status_code == 200
    after = [list(map(tuple, v)) for v in dump().values()]
    assert before == after
