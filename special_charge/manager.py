"""
SpecialChargeManager — 唯一資料存取入口（比照 database.py::Database 的
扁平風格，不做 Repository/Service/Adapter 分層）。

依 Architecture Reset 的憲法：
    特殊收費是一個獨立模組，只能唯讀讀取月費人員，不得修改任何既有月費
    或記帳功能，不得新增任何與記帳相關的架構。功能僅限活動、人員、項目、
    繳費狀態、歷史查詢、Excel 匯出。

本檔案完全不 import charge_management／payment_engine（那些設計已依
Architecture Reset 指示忽略），只依賴標準函式庫 sqlite3 與 openpyxl。
"""

import re
import sqlite3
from datetime import datetime, date, timedelta
from typing import Optional

from . import schema

# S15：繳收期限
DUE_DEFAULT_DAYS = 14        # 預計繳收完畢日期預設＝活動開始日期＋14 天（你的決定 Q9 B）
DUE_SOON_DAYS = 3            # 期限前 3 天內（含當天）且還有人未繳 →「即將到期」（Q11 B）
STATUS_TEXT = {"active": "進行中", "completed": "已完成"}     # S15（Q18 A）
_KEEP = object()


def _to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def default_due_date(start) -> str:
    """S15：預設的預計繳收完畢日期（開始日期＋14 天）；開始日期無效時用今天＋14 天。"""
    d = _to_date(start) or datetime.now().date()
    return (d + timedelta(days=DUE_DEFAULT_DAYS)).isoformat()


def check_due_date(start, due) -> None:
    """S15：期限不可早於開始日期（兩者都有填時才檢查）。"""
    s_, d_ = _to_date(start), _to_date(due)
    if due and d_ is None:
        raise ValueError("預計繳收完畢日期格式不正確")
    if s_ and d_ and d_ < s_:
        raise ValueError("預計繳收完畢日期不可早於活動開始日期")


def alert_label(alert: dict) -> str:
    """S15：活動列表「繳收期限」欄的文字（桌面、Web 共用）。"""
    st = alert.get("state")
    if st == "overdue":
        return f"⚠ 逾期 {alert['days']} 天"
    if st == "soon":
        return f"⏰ {DUE_SOON_DAYS} 天內到期"
    if st == "done":
        return "✅ 已收齊"
    return alert.get("due_date") or ""


def alert_banner(alert: dict) -> str:
    """S15：活動明細上方的提示文字（桌面、Web 共用）。沒有要提示時回傳空字串。"""
    st, due = alert.get("state"), alert.get("due_date")
    rest = f"尚有 {alert.get('unpaid', 0)} 人未繳、未收 ${alert.get('outstanding', 0):,.0f}"
    if st == "overdue":
        return f"⚠ 已超過繳收期限 {alert['days']} 天（{due}），{rest}"
    if st == "soon":
        left = alert["days_left"]
        when = "今天到期" if left == 0 else f"還剩 {left} 天"
        return f"⏰ 繳收期限 {due}，{when}，{rest}"
    if st == "normal":
        return f"繳收期限：{due}（還有 {alert['days_left']} 天）"
    if st == "done":
        return "✅ 已收齊" + (f"（繳收期限 {due}）" if due else "")
    return ""


class SpecialChargeManager:
    """特殊收費管理的唯一資料存取類別。"""

    # ★ 依你的最終要求：Excel「個人明細」Sheet 依年級排序，固定順序
    # 六→五→四→三→二→一，沒有對應年級文字的排在最後。
    _GRADE_ORDER = {"六年級": 6, "五年級": 5, "四年級": 4, "三年級": 3, "二年級": 2, "一年級": 1}

    # S11：臨時人員的身分（你的決定 Q7：就這 5 種）。排序在所有學生之後，依此順序。
    GUEST_ROLES = ("管理人員", "教練", "家長", "校外選手", "其他")

    def __init__(self, conn: sqlite3.Connection, ensure_schema: bool = True) -> None:
        self.conn = conn
        # ★ Web 唯讀重用（2026 新增，向下相容）：桌面版建構這個類別時一律
        # 執行 `CREATE TABLE IF NOT EXISTS`（idempotent，但仍是一次 DDL
        # 寫入）。Web 版是嚴格唯讀連線，不應該對連線送出任何 DDL，所以
        # 允許 `ensure_schema=False` 跳過這一步——桌面版所有既有呼叫端
        # 都沒有傳這個參數，預設值 `True` 維持完全相同的行為，不影響
        # 桌面版任何既有功能。
        if ensure_schema:
            schema.ensure_schema(self.conn)

    # ── 活動 ──────────────────────────────────────────────────────────────

    def create_activity(self, name: str, activity_date: str = "", note: str = "",
                        due_date: Optional[str] = None) -> int:
        """建立活動。★依你的裁定：活動名稱不要求唯一，用 id 區分即可。
        S15：activity_date＝活動開始日期；due_date＝預計繳收完畢日期（可空）。"""
        check_due_date(activity_date, due_date)
        cur = self.conn.execute(
            "INSERT INTO special_activity (name, activity_date, status, note, created_at, due_date) "
            "VALUES (?, ?, 'active', ?, ?, ?)",
            (name, activity_date, note, datetime.now().isoformat(timespec="seconds"), due_date or None),
        )
        self.conn.commit()
        return cur.lastrowid

    def list_activities(self, status: Optional[str] = None, year: Optional[str] = None) -> list[dict]:
        """列出活動，可依 status 篩選（'active'／'completed'／None=全部），
        ★ Beta P2：也可依年度篩選（`year` 取自 activity_date 的前 4 碼，
        沒有 activity_date 則退而取 created_at 的前 4 碼）。附帶每個活動
        的「已繳人數/總人數」，供主畫面左側列表顯示。"""
        if status:
            rows = self.conn.execute(
                "SELECT * FROM special_activity WHERE status=? ORDER BY id DESC", (status,)
            ).fetchall()
        else:
            rows = self.conn.execute(
                "SELECT * FROM special_activity ORDER BY id DESC"
            ).fetchall()

        result = []
        for row in rows:
            activity = dict(row)
            source_date = activity["activity_date"] or activity["created_at"] or ""
            activity["year"] = source_date[:4] if len(source_date) >= 4 else ""
            if year and activity["year"] != year:
                continue
            total = self.conn.execute(
                "SELECT COUNT(*) FROM special_activity_member WHERE activity_id=?",
                (activity["id"],),
            ).fetchone()[0]
            paid = self.conn.execute(
                "SELECT COUNT(*) FROM special_activity_member WHERE activity_id=? AND paid=1",
                (activity["id"],),
            ).fetchone()[0]
            activity["member_total"] = total
            activity["member_paid"] = paid
            activity.setdefault("due_date", None)        # S15：Web 讀到未升級的資料庫時沒有這個欄位
            activity["alert"] = self.activity_alert(activity)
            result.append(activity)
        return result

    # ── S15：繳收期限提示 ─────────────────────────────────────────────────

    def activity_alert(self, activity, today=None) -> dict:
        """S15：活動的繳收狀態（只讀）。activity：活動 id 或 dict。
        state：completed（已完成，不提示，Q12 A）／done（需要繳費的人都已繳＝已收齊）／
               none（沒有期限）／normal（距期限超過 3 天）／soon（0～3 天）／overdue（已過期限）。
        應收 0 元的人視為不需繳費。"""
        a = self.get_activity(activity) if isinstance(activity, int) else activity
        members = self.list_members(a["id"])
        need = [m for m in members if not m["paid"] and m["receivable"] > 0]
        outstanding = sum(m["receivable"] for m in members) - self.collected_of(members)
        due_s = a.get("due_date") or None
        due = _to_date(due_s)
        out = {"state": "none", "due_date": due_s if due else None, "days": 0, "days_left": None,
               "unpaid": len(need), "outstanding": max(outstanding, 0.0),
               "unpaid_ids": [m["id"] for m in need]}
        if a.get("status") == "completed":
            out["state"] = "completed"
            return out
        if members and not need:
            out["state"] = "done"
            return out
        if due is None:
            return out
        d = _to_date(today) or datetime.now().date()
        left = (due - d).days
        out["days_left"] = left
        if left < 0:
            out["state"], out["days"] = "overdue", -left
        elif left <= DUE_SOON_DAYS:
            out["state"] = "soon"
        else:
            out["state"] = "normal"
        return out

    def count_overdue_activities(self, today=None) -> int:
        """S15：逾期的進行中活動數（主視窗分頁「⚠數字」用）。"""
        return sum(1 for a in self.conn.execute(
            "SELECT * FROM special_activity WHERE status='active'").fetchall()
            if self.activity_alert(dict(a), today)["state"] == "overdue")

    def list_available_years(self) -> list[str]:
        """★ Beta P2：供 UI 年度篩選下拉選單使用，列出目前所有活動涵蓋的年度。"""
        rows = self.conn.execute("SELECT activity_date, created_at FROM special_activity").fetchall()
        years = set()
        for activity_date, created_at in rows:
            source = activity_date or created_at or ""
            if len(source) >= 4:
                years.add(source[:4])
        return sorted(years, reverse=True)

    def get_activity(self, activity_id: int) -> dict:
        row = self.conn.execute(
            "SELECT * FROM special_activity WHERE id=?", (activity_id,)
        ).fetchone()
        if row is None:
            raise KeyError(f"special_activity id={activity_id} 不存在")
        return dict(row)

    def update_activity(self, activity_id: int, name: str, activity_date: str = "", note: str = "",
                        due_date=_KEEP) -> None:
        """★ Beta P1：建立活動後仍可修改名稱/日期/備註（不需要重建整個活動）。
        S15：due_date 沒有傳入時不變；傳 None 或空字串＝清除期限。"""
        cur = self.get_activity(activity_id)
        due = cur.get("due_date") if due_date is _KEEP else (due_date or None)
        check_due_date(activity_date, due)
        self.conn.execute(
            "UPDATE special_activity SET name=?, activity_date=?, note=?, due_date=? WHERE id=?",
            (name, activity_date, note, due, activity_id),
        )
        self.conn.commit()

    def duplicate_activity(self, activity_id: int, new_name: str, new_activity_date: str = "") -> int:
        """★ Beta P3：複製活動——把所有參與人員與項目複製到新活動，
        繳費狀態重置（新活動是全新的一次收費，不該帶著舊活動的已繳紀錄）。
        """
        new_activity_id = self.create_activity(new_name, new_activity_date)
        members = self.list_members(activity_id)
        for m in members:
            if m.get("is_guest"):      # S11：臨時人員一併複製
                new_participant_id = self.conn.execute(
                    "INSERT INTO special_activity_member (activity_id, tenant_id, paid, guest_name, guest_role) "
                    "VALUES (?, NULL, 0, ?, ?)",
                    (new_activity_id, m["tenant_name"], m["grade"]),
                ).lastrowid
            else:
                new_participant_id = self.conn.execute(
                    "INSERT INTO special_activity_member (activity_id, tenant_id, paid) VALUES (?, ?, 0)",
                    (new_activity_id, m["tenant_id"]),
                ).lastrowid
            for item in self.list_items(m["id"]):
                self.conn.execute(
                    "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
                    "VALUES (?, ?, ?, ?)",
                    (new_participant_id, item["item_name"], item["unit_price"], item["quantity"]),
                )
        self.conn.commit()
        return new_activity_id

    def complete_activity(self, activity_id: int) -> None:
        """★依你的裁定：不強制全員已繳才能完成；呼叫端（UI）自行決定是否
        先跳提醒，本方法只負責真正切換狀態，不做業務判斷。"""
        self.get_activity(activity_id)  # 不存在時自然拋 KeyError
        self.conn.execute(
            "UPDATE special_activity SET status='completed', completed_at=? WHERE id=?",
            (datetime.now().isoformat(timespec="seconds"), activity_id),
        )
        self.conn.commit()

    def reopen_activity(self, activity_id: int) -> None:
        """把已完成的活動改回進行中（例如誤按完成）。"""
        self.get_activity(activity_id)
        self.conn.execute(
            "UPDATE special_activity SET status='active', completed_at=NULL WHERE id=?",
            (activity_id,),
        )
        self.conn.commit()

    def delete_activity(self, activity_id: int) -> None:
        """★ 依你的要求新增：刪除整個活動，連帶刪除所有參與人員與收費項目
        （活動資料/參與人員/收費項目/繳費狀態全部一起刪除，不留孤兒資料）。
        不影響 rent_tenants（唯讀參照，只刪自己這邊的關聯資料）。
        UI 呼叫本方法前必須先跳二次確認，本方法本身不再另外確認，
        呼叫了就直接執行（無法復原）。"""
        self.get_activity(activity_id)  # 不存在時自然拋 KeyError
        member_ids = [
            row[0] for row in self.conn.execute(
                "SELECT id FROM special_activity_member WHERE activity_id=?", (activity_id,)
            ).fetchall()
        ]
        for member_id in member_ids:
            self.conn.execute("DELETE FROM special_activity_item WHERE participant_id=?", (member_id,))
        self.conn.execute("DELETE FROM special_activity_member WHERE activity_id=?", (activity_id,))
        self.conn.execute("DELETE FROM special_activity WHERE id=?", (activity_id,))
        self.conn.commit()

    def count_unpaid_members(self, activity_id: int) -> int:
        """供 UI 在「移至歷史」前顯示「尚有 N 人未完成繳費」提醒用。"""
        return self.conn.execute(
            "SELECT COUNT(*) FROM special_activity_member WHERE activity_id=? AND paid=0",
            (activity_id,),
        ).fetchone()[0]

    # ── 參與人員 ──────────────────────────────────────────────────────────

    def load_active_tenants_as_members(self, activity_id: int, tenant_ids: list[int]) -> None:
        """依 UI Wizard Step2 使用者勾選的結果，批次加入參與人員。
        `tenant_ids` 由呼叫端（View）先查詢 `rent_tenants WHERE active=1`
        取得清單、經使用者勾選後傳入——本方法本身不查詢 rent_tenants，
        避免 Manager 與既有 `Database` 類別的查詢邏輯重複或耦合。"""
        for tenant_id in tenant_ids:
            self.conn.execute(
                "INSERT INTO special_activity_member (activity_id, tenant_id, paid) "
                "VALUES (?, ?, 0)",
                (activity_id, tenant_id),
            )
        self.conn.commit()

    def is_tenant_already_member(self, activity_id: int, tenant_id: int) -> bool:
        """★ 新增：確認某個月費人員是否已經是這個活動的參與者
        （用來阻止重複加入）。"""
        row = self.conn.execute(
            "SELECT 1 FROM special_activity_member WHERE activity_id=? AND tenant_id=?",
            (activity_id, tenant_id),
        ).fetchone()
        return row is not None

    def add_member(self, activity_id: int, tenant_id: int) -> int:
        """★ 新增：活動建立後，重新加入單一參與人員（例如誤刪或漏選後
        補回）。新加入的人：沒有任何收費項目，應收金額為 0，未繳。
        不可重複加入同一活動中已存在的人員——重複加入會拋出 ValueError，
        呼叫端（UI）應該先用 `is_tenant_already_member()` 檢查並擋下，
        這裡的例外是最後一道防線，不是唯一防線。
        新增這個人完全不影響活動中其他既有參與者的項目與繳費狀態
        （純粹是 INSERT 一筆新的 `special_activity_member`，不觸碰任何
        既有列）。"""
        self.get_activity(activity_id)  # 活動不存在時自然拋 KeyError
        if self.is_tenant_already_member(activity_id, tenant_id):
            raise ValueError(f"tenant_id={tenant_id} 已經是這個活動的參與者，不可重複加入")
        cur = self.conn.execute(
            "INSERT INTO special_activity_member (activity_id, tenant_id, paid) VALUES (?, ?, 0)",
            (activity_id, tenant_id),
        )
        self.conn.commit()
        return cur.lastrowid

    # ── S11：臨時人員 ──────────────────────────────────────────────────────

    def _check_guest(self, name: str, role: str):
        name = (name or "").strip()
        if not name:
            raise ValueError("臨時人員姓名不可為空")
        if role not in self.GUEST_ROLES:
            raise ValueError(f"身分必須是：{'、'.join(self.GUEST_ROLES)}")
        return name

    def add_guest(self, activity_id: int, name: str, role: str) -> int:
        """S11：加入臨時人員（不在人員名單的管理人員、教練、校外選手等）。"""
        self.get_activity(activity_id)
        name = self._check_guest(name, role)
        cur = self.conn.execute(
            "INSERT INTO special_activity_member (activity_id, tenant_id, paid, guest_name, guest_role) "
            "VALUES (?, NULL, 0, ?, ?)",
            (activity_id, name, role),
        )
        self.conn.commit()
        return cur.lastrowid

    def update_guest(self, participant_id: int, name: str, role: str) -> None:
        """S11：修改臨時人員的姓名與身分（人員名單裡的人不能用這個方法）。"""
        member = self.get_member(participant_id)
        if not member["is_guest"]:
            raise ValueError("只有臨時人員可以修改姓名與身分")
        name = self._check_guest(name, role)
        self.conn.execute(
            "UPDATE special_activity_member SET guest_name=?, guest_role=? WHERE id=?",
            (name, role, participant_id),
        )
        self.conn.commit()

    def find_tenants_named(self, name: str) -> list[dict]:
        """S11：人員名單中同名的人（加入臨時人員前提醒用），含已停用。"""
        rows = self.conn.execute(
            "SELECT id, name, grade, active FROM rent_tenants WHERE name=?", ((name or "").strip(),)
        ).fetchall()
        return [dict(r) for r in rows]

    def list_eligible_tenants_to_add(self, activity_id: int, include_inactive: bool = False) -> list[dict]:
        """★ 新增：列出目前月費有效人員中，尚未參加這個活動的人
        （供 UI「＋ 新增參與人員」對話框使用）。

        S9：`include_inactive=True` 時也列出已停用的人（例：最後收費月份已過、
        但仍參加活動的人）；預設 False，結果與原本相同。附帶 `active` 欄位。
        排序改為年級（六→一）再依姓名，與其他畫面一致。"""
        sql = ("SELECT id, name, grade, active FROM rent_tenants "
               "WHERE id NOT IN (SELECT tenant_id FROM special_activity_member WHERE activity_id=?)")
        if not include_inactive:
            sql += " AND active=1"
        rows = self.conn.execute(sql, (activity_id,)).fetchall()
        result = [{"id": r["id"], "name": r["name"], "grade": r["grade"] or "",
                   "active": int(r["active"])} for r in rows]
        return self.sort_tenants_by_grade(result)

    def list_tenants_for_new_activity(self, include_inactive: bool = False) -> list[dict]:
        """S9：建立活動精靈第 2 步的人員清單（年級六→一、同年級依姓名）。"""
        sql = "SELECT id, name, grade, active FROM rent_tenants"
        if not include_inactive:
            sql += " WHERE active=1"
        rows = self.conn.execute(sql).fetchall()
        result = [{"id": r["id"], "name": r["name"], "grade": r["grade"] or "",
                   "active": int(r["active"])} for r in rows]
        return self.sort_tenants_by_grade(result)

    def sort_tenants_by_grade(self, tenants: list[dict]) -> list[dict]:
        """S9：人員清單（dict 含 name／grade）依年級六→一、同年級依姓名排序。
        與 `sort_members_by_grade` 使用同一份 `_GRADE_ORDER`。"""
        return sorted(tenants, key=lambda t: (-self._GRADE_ORDER.get(t["grade"], 0), t["name"]))

    def sort_members_by_grade(self, members: list[dict]) -> list[dict]:
        """★ 唯一的年級排序邏輯來源（Excel「個人明細」與 Web 查詢明細頁面
        皆呼叫這個方法，不各自重寫一份排序 key，避免兩處邏輯不一致）。

        依年級排序 六→五→四→三→二→一，沒有對應年級文字的排在最後；
        同年級再依姓名排序。排序鍵直接讀取 `member["grade"]`（唯讀
        JOIN `rent_tenants.grade` 取得的實際資料），不是用
        `id`/`participant_id`/`tenant_id`/建立順序或姓名反推年級。
        """
        return sorted(members, key=self.member_sort_key)

    def sort_members_done_last(self, members: list[dict]) -> list[dict]:
        """S15（Q14 A）：人員表「未繳在上、已繳在下」，兩區各自依 member_sort_key（年級、臨時人員在後）。
        桌面人員表與 Web 明細頁共用；表格輸入與 Excel 仍用 sort_members_by_grade（不變）。"""
        return sorted(members, key=lambda m: (1 if m["paid"] else 0, self.member_sort_key(m)))

    def member_sort_key(self, m: dict):
        """S11：學生依年級（六→一）、姓名；臨時人員排在所有學生之後，
        依身分（GUEST_ROLES 的順序）、姓名。桌面人員清單、表格輸入、Excel、Web 共用。"""
        if m.get("is_guest"):
            role = m.get("grade") or ""
            if role in self._GRADE_ORDER:
                # S13：結轉時刪除的停用人員改成臨時人員、身分欄是原本的年級 → 仍依年級排在學生之間
                return (0, -self._GRADE_ORDER[role], 0, m["tenant_name"])
            ridx = self.GUEST_ROLES.index(role) if role in self.GUEST_ROLES else len(self.GUEST_ROLES)
            return (1, 0, ridx, m["tenant_name"])
        return (0, -self._GRADE_ORDER.get(m["grade"], 0), 0, m["tenant_name"])

    def list_members(self, activity_id: int) -> list[dict]:
        """列出活動的參與人員，附帶即時計算的應收金額（不落地儲存），
        以及唯讀查詢來的姓名/年級（★ Beta P2：年級欄，若 rent_tenants
        沒有年級資料則留空，不強制要求）。"""
        rows = self.conn.execute(
            "SELECT * FROM special_activity_member WHERE activity_id=? ORDER BY id",
            (activity_id,),
        ).fetchall()
        result = []
        for row in rows:
            member = self._decorate_member(dict(row))
            member["receivable"] = self.get_member_receivable(member["id"])
            result.append(member)
        return result

    def _decorate_member(self, member: dict) -> dict:
        """加上 tenant_name／grade／is_guest。S11：臨時人員（tenant_id 為空）的姓名取
        guest_name、年級欄顯示身分 guest_role。Web 可能讀到尚未升級的舊資料庫（沒有
        guest 欄位），此時全部視為人員名單裡的人，與原本相同。"""
        if member.get("tenant_id") is None and member.get("guest_name"):
            member["is_guest"] = True
            member["tenant_name"] = member["guest_name"]
            member["grade"] = member.get("guest_role") or ""
            return member
        member["is_guest"] = False
        tenant_row = self.conn.execute(
            "SELECT name, grade FROM rent_tenants WHERE id=?", (member["tenant_id"],)
        ).fetchone()
        member["tenant_name"] = tenant_row["name"] if tenant_row else f"(未知會員 #{member['tenant_id']})"
        member["grade"] = (tenant_row["grade"] if tenant_row and tenant_row["grade"] else "") or ""
        return member

    def get_member(self, participant_id: int) -> dict:
        row = self.conn.execute(
            "SELECT * FROM special_activity_member WHERE id=?", (participant_id,)
        ).fetchone()
        if row is None:
            raise KeyError(f"special_activity_member id={participant_id} 不存在")
        return self._decorate_member(dict(row))

    def count_items(self, participant_id: int) -> int:
        """供 UI 顯示移除確認警告用（「共有 N 筆項目，確定移除？」）。"""
        return self.conn.execute(
            "SELECT COUNT(*) FROM special_activity_item WHERE participant_id=?",
            (participant_id,),
        ).fetchone()[0]

    def remove_member(self, participant_id: int) -> None:
        """★依你的裁定：允許移除，即使已有項目資料——連帶刪除其項目。
        UI 呼叫本方法前，必須先用 `count_items()` 查詢並顯示確認警告，
        本方法本身不跳警告（那是 UI 層責任），呼叫了就直接執行。"""
        self.get_member(participant_id)  # 不存在時自然拋 KeyError
        self.conn.execute(
            "DELETE FROM special_activity_item WHERE participant_id=?", (participant_id,)
        )
        self.conn.execute(
            "DELETE FROM special_activity_member WHERE id=?", (participant_id,)
        )
        self.conn.commit()

    # ── 收費項目 ──────────────────────────────────────────────────────────

    def add_item(self, participant_id: int, item_name: str, unit_price: float, quantity: int = 1) -> int:
        cur = self.conn.execute(
            "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
            "VALUES (?, ?, ?, ?)",
            (participant_id, item_name, unit_price, quantity),
        )
        self.conn.commit()
        return cur.lastrowid

    def update_item(self, item_id: int, item_name: str, unit_price: float, quantity: int) -> None:
        """★ Beta P1：數量改為 0 視同不需要該項目，直接刪除，不留
        quantity=0 的殘留資料列。"""
        if quantity <= 0:
            self.remove_item(item_id)
            return
        self.conn.execute(
            "UPDATE special_activity_item SET item_name=?, unit_price=?, quantity=? WHERE id=?",
            (item_name, unit_price, quantity, item_id),
        )
        self.conn.commit()

    def remove_item(self, item_id: int) -> None:
        self.conn.execute("DELETE FROM special_activity_item WHERE id=?", (item_id,))
        self.conn.commit()

    def list_items(self, participant_id: int) -> list[dict]:
        rows = self.conn.execute(
            "SELECT * FROM special_activity_item WHERE participant_id=? ORDER BY id",
            (participant_id,),
        ).fetchall()
        result = []
        for row in rows:
            item = dict(row)
            item["subtotal"] = item["unit_price"] * item["quantity"]  # ★不落地儲存，即時算
            result.append(item)
        return result

    def apply_item_to_all_members(self, activity_id: int, item_name: str, unit_price: float, quantity: int = 1) -> None:
        """★依你的建議新增：「全部套用」——一次幫該活動所有參與人員都加上
        同一個項目，之後使用者再對少數人個別調整/刪除，避免 30 人逐一輸入
        同樣項目的重複工作。"""
        members = self.conn.execute(
            "SELECT id FROM special_activity_member WHERE activity_id=?", (activity_id,)
        ).fetchall()
        for (participant_id,) in members:
            self.conn.execute(
                "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
                "VALUES (?, ?, ?, ?)",
                (participant_id, item_name, unit_price, quantity),
            )
        self.conn.commit()

    def preview_apply_item(self, participant_ids: list[int], item_name: str) -> dict:
        """S9：批次新增項目前的預覽——多少人會新增、多少人已有同名項目（會被改成
        新的單價與數量）、其中多少人已繳（應收金額會改變）。"""
        add, update, paid_names = 0, 0, []
        for pid in participant_ids:
            member = self.get_member(pid)
            exists = self.conn.execute(
                "SELECT 1 FROM special_activity_item WHERE participant_id=? AND item_name=?",
                (pid, item_name),
            ).fetchone()
            if exists:
                update += 1
            else:
                add += 1
            if member["paid"]:
                paid_names.append(member["tenant_name"])
        return {"add": add, "update": update, "paid_names": paid_names}

    def apply_item_to_members(self, participant_ids: list[int], item_name: str,
                              unit_price: float, quantity: int = 1) -> dict:
        """S9：對指定的參與人員套用同一個項目（全部／某年級／勾選的人由 UI 決定名單）。

        依你的決定（B2-A）：已有同名項目的人 → 改成新的單價與數量（同名有多筆時
        只改第一筆）；沒有的人 → 新增一筆。同一個 transaction 內完成。
        既有的 `apply_item_to_all_members` 不變（行為：一律新增）。
        """
        added, updated = 0, 0
        try:
            for pid in participant_ids:
                row = self.conn.execute(
                    "SELECT id FROM special_activity_item WHERE participant_id=? AND item_name=? "
                    "ORDER BY id LIMIT 1",
                    (pid, item_name),
                ).fetchone()
                if row:
                    self.conn.execute(
                        "UPDATE special_activity_item SET unit_price=?, quantity=? WHERE id=?",
                        (unit_price, quantity, row[0]),
                    )
                    updated += 1
                else:
                    self.conn.execute(
                        "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
                        "VALUES (?, ?, ?, ?)",
                        (pid, item_name, unit_price, quantity),
                    )
                    added += 1
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return {"added": added, "updated": updated}

    def mark_paid_many(self, participant_ids: list[int], paid_date: str) -> dict:
        """S9：批次標記已繳。繳費金額＝各自目前的應收；已經是已繳的人略過（不改日期）。
        同一個 transaction 內完成。"""
        marked, skipped = 0, 0
        try:
            for pid in participant_ids:
                member = self.get_member(pid)
                if member["paid"]:
                    skipped += 1
                    continue
                self.conn.execute(
                    "UPDATE special_activity_member SET paid=1, paid_amount=?, paid_date=? WHERE id=?",
                    (self.get_member_receivable(pid), paid_date, pid),
                )
                marked += 1
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return {"marked": marked, "skipped": skipped}

    # ── S11：價目表（由現有項目推算，不另外保存；你的決定 Q4）─────────────────

    def price_list(self, activity_id: int) -> list[dict]:
        """活動目前所有項目名稱，每個名稱的欄位單價＝最多人使用的單價
        （同數時取最近建立的那一筆的單價）。依項目第一次建立的順序排列。
        回傳 [{"name", "price", "count", "cells"}]（count＝使用這個單價的格數，cells＝總格數）。"""
        rows = self.conn.execute(
            "SELECT i.id, i.item_name, i.unit_price FROM special_activity_item i "
            "JOIN special_activity_member m ON m.id = i.participant_id "
            "WHERE m.activity_id=? ORDER BY i.id",
            (activity_id,),
        ).fetchall()
        info: dict = {}
        for item_id, name, price in rows:
            d = info.setdefault(name, {"first": item_id, "prices": {}, "cells": 0})
            pc = d["prices"].setdefault(price, [0, 0])
            pc[0] += 1
            pc[1] = item_id
            d["cells"] += 1
        out = []
        for name, d in sorted(info.items(), key=lambda kv: kv[1]["first"]):
            price, (count, _) = max(d["prices"].items(), key=lambda kv: (kv[1][0], kv[1][1]))
            out.append({"name": name, "price": price, "count": count, "cells": d["cells"]})
        return out

    def grid_cells(self, activity_id: int) -> dict:
        """表格輸入用：{participant_id: {item_name: {"item_id", "price", "qty"} 或 {"multi": True}}}。
        同一人有兩筆以上同名項目時標示 multi（表格鎖定，只能在個人項目視窗修改）。"""
        rows = self.conn.execute(
            "SELECT i.id, i.participant_id, i.item_name, i.unit_price, i.quantity "
            "FROM special_activity_item i JOIN special_activity_member m ON m.id = i.participant_id "
            "WHERE m.activity_id=? ORDER BY i.id",
            (activity_id,),
        ).fetchall()
        cells: dict = {}
        for item_id, pid, name, price, qty in rows:
            person = cells.setdefault(pid, {})
            if name in person:
                prev = person[name]
                person[name] = {"multi": True,
                                "subtotal": prev.get("subtotal", prev.get("price", 0) * prev.get("qty", 0)) + price * qty}
            else:
                person[name] = {"item_id": item_id, "price": price, "qty": qty}
        return cells

    def save_grid(self, activity_id: int, cells: list[dict], update_paid_ids=()) -> dict:
        """S11：表格輸入一次儲存（同一個 transaction）。

        cells：要寫入的格子 [{"participant_id", "item_name", "price", "qty"}]（只需傳有變動的格子）。
          qty <= 0 → 刪除這個人的這個項目；已有 → 改單價與數量；沒有 → 新增。
          同一人有多筆同名項目（multi）的格子不處理（表格不允許改）。
        update_paid_ids：已繳的人中，要把繳費金額更新為新應收的人（你的決定 Q5 A）；
          其餘已繳的人繳費金額不變。
        回傳 {"added", "updated", "deleted", "paid_updated"}。"""
        member_ids = {r[0] for r in self.conn.execute(
            "SELECT id FROM special_activity_member WHERE activity_id=?", (activity_id,))}
        counts = {"added": 0, "updated": 0, "deleted": 0, "paid_updated": 0}
        self.conn.commit()
        try:
            for c in cells:
                pid, name = c["participant_id"], c["item_name"]
                if pid not in member_ids:
                    raise ValueError(f"參與者 {pid} 不在這個活動")
                qty = int(c.get("qty") or 0)
                price = float(c.get("price") or 0)
                existing = self.conn.execute(
                    "SELECT id, unit_price, quantity FROM special_activity_item "
                    "WHERE participant_id=? AND item_name=?", (pid, name)).fetchall()
                if len(existing) > 1:
                    continue
                if qty <= 0:
                    if existing:
                        self.conn.execute("DELETE FROM special_activity_item WHERE id=?", (existing[0][0],))
                        counts["deleted"] += 1
                    continue
                if price <= 0:
                    raise ValueError(f"「{name}」單價必須大於 0")
                if existing:
                    item_id, old_price, old_qty = existing[0]
                    if old_price != price or old_qty != qty:
                        self.conn.execute(
                            "UPDATE special_activity_item SET unit_price=?, quantity=? WHERE id=?",
                            (price, qty, item_id))
                        counts["updated"] += 1
                else:
                    self.conn.execute(
                        "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
                        "VALUES (?, ?, ?, ?)", (pid, name, price, qty))
                    counts["added"] += 1
            for pid in update_paid_ids:
                if pid in member_ids:
                    cur = self.conn.execute(
                        "UPDATE special_activity_member SET paid_amount=? WHERE id=? AND paid=1",
                        (self.get_member_receivable(pid), pid))
                    counts["paid_updated"] += cur.rowcount
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return counts

    def paid_amount_changes(self, participant_ids, new_receivables: Optional[dict] = None) -> list[dict]:
        """S11（Q5）：已繳的人中，應收和繳費金額不同的人。
        new_receivables：{participant_id: 新應收}（表格尚未儲存時由畫面提供）；沒有則用目前應收。
        回傳 [{"id", "name", "paid_amount", "receivable"}]。"""
        out = []
        for pid in participant_ids:
            m = self.get_member(pid)
            if not m["paid"]:
                continue
            rec = (new_receivables or {}).get(pid, self.get_member_receivable(pid))
            paid = m["paid_amount"] if m["paid_amount"] is not None else rec
            if abs(rec - paid) > 1e-9:
                out.append({"id": pid, "name": m["tenant_name"], "paid_amount": paid, "receivable": rec})
        return out

    def sync_paid_amount(self, participant_ids) -> int:
        """S11（Q5「更新」）：已繳的人繳費金額改為目前應收（繳費日期不變）。"""
        n = 0
        try:
            for pid in participant_ids:
                cur = self.conn.execute(
                    "UPDATE special_activity_member SET paid_amount=? WHERE id=? AND paid=1",
                    (self.get_member_receivable(pid), pid))
                n += cur.rowcount
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return n

    @staticmethod
    def collected_of(members: list[dict]) -> float:
        """S11（定案）：已收＝已繳的人的繳費金額合計（Web 原本就是這樣算）。
        舊資料若已繳但沒有繳費金額，以應收計。"""
        return sum((m["paid_amount"] if m["paid_amount"] is not None else m["receivable"])
                   for m in members if m["paid"])

    def get_activity_collected(self, activity_id: int) -> float:
        return self.collected_of(self.list_members(activity_id))

    # ── S11：項目名稱提示（以前用過的名稱與最近單價）──────────────────────────

    def item_name_history(self) -> list[dict]:
        """所有活動用過的項目名稱，各自附「最近一次」的單價與活動。
        最近＝活動日期最新（沒有日期用建立時間），同一活動取最後建立的項目。
        回傳依名稱排序 [{"name", "price", "activity", "date"}]。只讀，不另外保存。"""
        rows = self.conn.execute(
            "SELECT i.item_name, i.unit_price, i.id, a.name, "
            "COALESCE(NULLIF(a.activity_date, ''), substr(a.created_at, 1, 10)) AS d "
            "FROM special_activity_item i "
            "JOIN special_activity_member m ON m.id = i.participant_id "
            "JOIN special_activity a ON a.id = m.activity_id"
        ).fetchall()
        latest: dict = {}
        for name, price, item_id, act, d in rows:
            key = (d or "", item_id)
            if name not in latest or key > latest[name][0]:
                latest[name] = (key, price, act, d or "")
        return [{"name": n, "price": v[1], "activity": v[2], "date": v[3]}
                for n, v in sorted(latest.items())]

    # ── S11.1：以舊活動為範本建立（取代「複製活動」）──────────────────────────

    def template_people(self, source_id: int, include_inactive: bool = False) -> dict:
        """範本第 3 步的人員清單。
        tenants：目前人員（依年級排序；include_inactive 時含已停用），in_source＝上次有參加。
        guests：來源活動的臨時人員（你的決定 Q8：列出並預先勾選）。
        inactive_in_source：上次有參加、現在已停用的人（不預先勾選，畫面上提示）。"""
        src = self.list_members(source_id)
        src_tids = {m["tenant_id"] for m in src if not m["is_guest"]}
        tenants = self.list_tenants_for_new_activity(include_inactive=include_inactive)
        for t in tenants:
            t["in_source"] = t["id"] in src_tids
        def is_inactive(tid):
            row = self.conn.execute("SELECT active FROM rent_tenants WHERE id=?", (tid,)).fetchone()
            return row is not None and row[0] == 0
        inactive = [m["tenant_name"] for m in self.sort_members_by_grade(src)
                    if not m["is_guest"] and is_inactive(m["tenant_id"])]
        guests = [{"participant_id": m["id"], "name": m["tenant_name"], "role": m["grade"]}
                  for m in self.sort_members_by_grade(src) if m["is_guest"]]
        return {"tenants": tenants, "guests": guests, "inactive_in_source": inactive}

    def create_from_template(self, source_id: int, name: str, activity_date: str,
                             prices: list, tenant_ids: list, source_guest_ids: list = (),
                             new_guests: list = (), carry_quantities: bool = True,
                             due_date: Optional[str] = None) -> dict:
        """S11.1：以來源活動為範本建立新活動（同一個 transaction）。

        prices：新的價目表 [{"name", "price"}]（第 2 步可改單價、刪除、新增）。
        tenant_ids：勾選的人員名單成員；source_guest_ids：勾選的來源臨時人員（參與者 id）；
        new_guests：新加的臨時人員 [(姓名, 身分)]。
        carry_quantities（Q3 A）：帶入上次的數量——只帶價目表中仍存在的項目；
          來源中單價＝來源欄位單價的格子改用新單價；個別單價的格子保留原單價（Q4 A），列在 kept_overrides。
        繳費狀態、日期、備註都不帶入。
        回傳 {"activity_id", "carried", "kept_overrides": [{"name", "item", "price", "column_price"}]}。"""
        name = (name or "").strip()
        if not name:
            raise ValueError("活動名稱不可為空")
        self.get_activity(source_id)
        check_due_date(activity_date, due_date)          # S15：期限不沿用來源活動，由畫面帶入（預設開始＋14 天）
        new_price = {}
        for p in prices:
            if not (p["name"] or "").strip() or float(p["price"]) <= 0:
                raise ValueError("價目表的項目名稱與單價皆須有效")
            new_price[p["name"].strip()] = float(p["price"])
        old_col = {p["name"]: p["price"] for p in self.price_list(source_id)}
        src_members = self.list_members(source_id)
        by_tid = {m["tenant_id"]: m for m in src_members if not m["is_guest"]}
        by_pid = {m["id"]: m for m in src_members}
        out = {"carried": 0, "kept_overrides": []}
        self.conn.commit()
        try:
            cur = self.conn.execute(
                "INSERT INTO special_activity (name, activity_date, status, note, created_at, due_date) "
                "VALUES (?, ?, 'active', '', ?, ?)",
                (name, activity_date, datetime.now().isoformat(timespec="seconds"), due_date or None))
            aid = cur.lastrowid
            plan = []                                     # (新參與者 id, 來源參與者 or None)
            for tid in tenant_ids:
                pid = self.conn.execute(
                    "INSERT INTO special_activity_member (activity_id, tenant_id, paid) VALUES (?, ?, 0)",
                    (aid, tid)).lastrowid
                plan.append((pid, by_tid.get(tid)))
            for spid in source_guest_ids:
                g = by_pid[spid]
                if not g["is_guest"]:
                    raise ValueError(f"參與者 {spid} 不是臨時人員")
                pid = self.conn.execute(
                    "INSERT INTO special_activity_member (activity_id, tenant_id, paid, guest_name, guest_role) "
                    "VALUES (?, NULL, 0, ?, ?)", (aid, g["tenant_name"], g["grade"])).lastrowid
                plan.append((pid, g))
            for gname, grole in new_guests:
                gname = self._check_guest(gname, grole)
                self.conn.execute(
                    "INSERT INTO special_activity_member (activity_id, tenant_id, paid, guest_name, guest_role) "
                    "VALUES (?, NULL, 0, ?, ?)", (aid, gname, grole))
            if carry_quantities:
                for pid, src in plan:
                    if src is None:
                        continue
                    for it in self.list_items(src["id"]):
                        nm = it["item_name"]
                        if nm not in new_price:
                            continue
                        price = new_price[nm]
                        if it["unit_price"] != old_col.get(nm):
                            price = it["unit_price"]
                            out["kept_overrides"].append({"name": src["tenant_name"], "item": nm,
                                                          "price": price, "column_price": new_price[nm]})
                        self.conn.execute(
                            "INSERT INTO special_activity_item (participant_id, item_name, unit_price, quantity) "
                            "VALUES (?, ?, ?, ?)", (pid, nm, price, it["quantity"]))
                        out["carried"] += 1
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        out["activity_id"] = aid
        return out

    # ── S11.1：產生記帳備註（你的決定 Q5 A、Q6 A：只產生文字，不連動記帳）─────────

    _CODE_KEY = re.compile(r"(?:後五碼|後5碼|末五碼|末5碼)[：:\s]*((?:\d{5}(?:\s*[、,，/／及和或\s]\s*)?)+)")

    @classmethod
    def parse_account_codes(cls, note: str) -> list:
        """從人員備註取出後五碼（可能 1～3 個帳號）。例：「後五碼：05843、中租」→ ["05843"]；
        「後五碼：12345、67890」→ ["12345", "67890"]。只認「後五碼／後5碼／末五碼」後面的 5 位數字。"""
        codes = []
        for m in cls._CODE_KEY.finditer(note or ""):
            for c in re.findall(r"\d{5}", m.group(1)):
                if c not in codes:
                    codes.append(c)
        return codes

    def note_text(self, activity_id: int, participant_ids: list, with_codes: bool = True) -> dict:
        """格式（Q6 A）：「活動名，姓名金額(後五碼)、姓名金額(後五碼)」，順序同人員清單（年級排序）。
        金額：已繳的人用繳費金額，未繳的人用應收。後五碼取自人員備註；臨時人員與沒有後五碼的人只寫姓名金額；
        有多個帳號的人全部列出（12345/67890），由你保留實際匯款的那個。
        回傳 {"text", "total", "count", "multi": [姓名], "missing": [姓名]}。"""
        activity = self.get_activity(activity_id)
        wanted = set(participant_ids)
        members = [m for m in self.sort_members_by_grade(self.list_members(activity_id)) if m["id"] in wanted]
        parts, total, multi, missing = [], 0.0, [], []
        for m in members:
            amount = m["paid_amount"] if (m["paid"] and m["paid_amount"] is not None) else m["receivable"]
            total += amount
            piece = f"{m['tenant_name']}{amount:.0f}"
            if with_codes and not m["is_guest"]:
                row = self.conn.execute("SELECT note FROM rent_tenants WHERE id=?", (m["tenant_id"],)).fetchone()
                codes = self.parse_account_codes(row[0] if row else "")
                if codes:
                    piece += f"({'/'.join(codes)})"
                    if len(codes) > 1:
                        multi.append(m["tenant_name"])
                else:
                    missing.append(m["tenant_name"])
            parts.append(piece)
        text = f"{activity['name']}，" + "、".join(parts) if parts else ""
        return {"text": text, "total": total, "count": len(members), "multi": multi, "missing": missing}

    # ── 即時計算（不落地儲存）──────────────────────────────────────────────

    def get_member_receivable(self, participant_id: int) -> float:
        row = self.conn.execute(
            "SELECT unit_price, quantity FROM special_activity_item WHERE participant_id=?",
            (participant_id,),
        ).fetchall()
        return sum(unit_price * quantity for unit_price, quantity in row)

    def get_activity_total_receivable(self, activity_id: int) -> float:
        members = self.conn.execute(
            "SELECT id FROM special_activity_member WHERE activity_id=?", (activity_id,)
        ).fetchall()
        return sum(self.get_member_receivable(m_id) for (m_id,) in members)

    # ── 繳費狀態 ──────────────────────────────────────────────────────────

    def mark_paid(self, participant_id: int, paid_amount: Optional[float] = None,
                  paid_date: Optional[str] = None, remark: str = "") -> None:
        """標記已繳。`paid_amount` 若未提供，預設帶入目前應收金額
        （使用者仍可在 UI 自行調整後再傳入不同的值，例如找零情境）。"""
        self.get_member(participant_id)
        if paid_amount is None:
            paid_amount = self.get_member_receivable(participant_id)
        if paid_date is None:
            paid_date = datetime.now().strftime("%Y-%m-%d")
        self.conn.execute(
            "UPDATE special_activity_member SET paid=1, paid_amount=?, paid_date=?, remark=? "
            "WHERE id=?",
            (paid_amount, paid_date, remark, participant_id),
        )
        self.conn.commit()

    def mark_unpaid(self, participant_id: int) -> None:
        self.get_member(participant_id)
        self.conn.execute(
            "UPDATE special_activity_member SET paid=0, paid_amount=NULL, paid_date=NULL WHERE id=?",
            (participant_id,),
        )
        self.conn.commit()

    def update_remark(self, participant_id: int, remark: str) -> None:
        self.get_member(participant_id)
        self.conn.execute(
            "UPDATE special_activity_member SET remark=? WHERE id=?", (remark, participant_id)
        )
        self.conn.commit()

    # ── Excel 匯出（固定兩個 Sheet，依你的規格）───────────────────────────────

    def export_excel(self, activity_id: int, path: Optional[str] = None):
        """匯出 Excel。★ Web 唯讀重用（2026 新增，向下相容）：`path` 改為
        可選——桌面版既有呼叫方式（傳入路徑）完全不變，會寫入該路徑並
        回傳 `None`；若不傳 `path`（`None`），改成回傳 `bytes`（不落地寫
        檔案到伺服器磁碟），供 Web 端 `StreamingResponse` 直接下載使用，
        不需要在伺服器暫存任何檔案。"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        activity = self.get_activity(activity_id)
        # ★ 需求變更（你這輪明確裁定）：「活動總表」與「個人明細」的人員
        # 排序規則必須完全一致——都依 sort_members_by_grade()（第一鍵年級
        # 六→五→四→三→二→一，第二鍵姓名）。不再保留「活動總表維持原始
        # 加入順序」的舊規則。
        members = self.sort_members_by_grade(self.list_members(activity_id))
        paid_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")

        wb = openpyxl.Workbook()

        # Sheet1：活動總表（★ Beta P2：年級欄／已繳用 V／已繳整列淺灰反色）
        ws1 = wb.active
        ws1.title = "活動總表"
        ws1.append(["年級", "姓名", "應收", "已繳", "日期"])
        for cell in ws1[1]:
            cell.font = Font(bold=True)
        for m in members:
            ws1.append([
                m["grade"],
                m["tenant_name"],
                m["receivable"],
                "V" if m["paid"] else "",
                m["paid_date"] or "",
            ])
            if m["paid"]:
                for cell in ws1[ws1.max_row]:
                    cell.fill = paid_fill
        for col in ws1.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws1.column_dimensions[col[0].column_letter].width = max_len + 4

        # ★ 依你的要求：Sheet1 最下面自動加統計摘要（不用另外做 Dashboard）
        total_count = len(members)
        paid_count = sum(1 for m in members if m["paid"])
        unpaid_count = total_count - paid_count
        receivable_total = sum(m["receivable"] for m in members)
        collected_total = self.collected_of(members)   # S11：已收＝繳費金額合計（與 Web 一致）
        outstanding_total = receivable_total - collected_total

        ws1.append([])
        summary_rows = [
            ("總人數", total_count),
            ("已繳", paid_count),
            ("未繳", unpaid_count),
            ("應收", receivable_total),
            ("已收", collected_total),
            ("未收", outstanding_total),
        ]
        for label, value in summary_rows:
            ws1.append([label, value])
            ws1.cell(row=ws1.max_row, column=1).font = Font(bold=True)

        # Sheet2：個人明細（★ 依你的最終要求：依年級排序 六→五→四→三→二→一，
        # 同年級再依姓名排序；排序鍵直接讀取 member 的 "grade" 欄位——這是
        # `list_members()` 唯讀 JOIN rent_tenants.grade 取得的實際資料，
        # 不是用 tenant_id、建立順序或姓名反推年級）
        # members 已經是 sort_members_by_grade() 排序過的結果（與 Sheet1
        # 共用同一份，確保兩個 Sheet 排序規則完全一致，不再各自排序一次）
        ws2 = wb.create_sheet("個人明細")
        row_idx = 1
        for m in members:
            # ★ 修正：先前姓名列沒有顯示年級，導致打開檔案完全看不到排序
            # 依據是什麼——只看得到姓名順序，看不出「年級」欄位本身。
            # 這裡把年級顯示在姓名前面，讓排序依據在檔案裡直接可見可驗證。
            label = f"{m['grade']}　{m['tenant_name']}" if m["grade"] else m["tenant_name"]
            ws2.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            row_idx += 1
            items = self.list_items(m["id"])
            for item in items:
                ws2.append([item["item_name"], item["unit_price"], item["quantity"], item["subtotal"]])
                row_idx += 1
            ws2.append(["合計", "", "", m["receivable"]])
            row_idx += 2  # 空一行分隔

        if path is None:
            import io
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return buf.read()

        wb.save(path)
        return None

    def _get_tenant_name(self, tenant_id: int) -> str:
        row = self.conn.execute(
            "SELECT name FROM rent_tenants WHERE id=?", (tenant_id,)
        ).fetchone()
        return row[0] if row else f"(未知會員 #{tenant_id})"
