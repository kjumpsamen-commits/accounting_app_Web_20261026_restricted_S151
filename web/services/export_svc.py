"""
services/export_svc.py
──────────────────────
Excel 匯出服務（唯讀）。

只讀取資料，產生 bytes 串流，不寫入任何檔案到伺服器磁碟。
邏輯參考桌面版 export_view.py，但不修改原始碼。
"""

import io
from datetime import date as _date

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from web.utils.db import ReadOnlyDB


# ── 樣式常數 ──────────────────────────────────────────────────────────────────

_HDR_FILL    = PatternFill("solid", fgColor="1E3A5F")
_HDR_FONT    = Font(bold=True, color="FFFFFF", size=11)
_INC_FILL    = PatternFill("solid", fgColor="ECFDF5")
_EXP_FILL    = PatternFill("solid", fgColor="FEF2F2")
_INC_FONT    = Font(color="065F46")
_EXP_FONT    = Font(color="991B1B")
_TOTAL_FILL  = PatternFill("solid", fgColor="1E3A5F")
_TOTAL_FONT  = Font(bold=True, color="FFFFFF")
_BORDER      = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center")


def _hdr(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill      = _HDR_FILL
        cell.font      = _HDR_FONT
        cell.alignment = _CENTER
        cell.border    = _BORDER


def _autowidth(ws):
    for col in ws.columns:
        max_len   = max((len(str(cell.value or "")) for cell in col), default=0)
        col_ltr   = get_column_letter(col[0].column)
        ws.column_dimensions[col_ltr].width = min(max_len + 4, 40)


# ── 主要匯出函式 ──────────────────────────────────────────────────────────────

def export_transactions_xlsx(
    db: ReadOnlyDB,
    date_from: str = "",
    date_to:   str = "",
    type_:     str = "",
    year_filter: str = "",
    include_detail:   bool = True,
    include_year:     bool = True,
    include_month:    bool = True,
    include_category: bool = True,
) -> bytes:
    """
    產生收支明細 Excel，回傳 bytes（供 StreamingResponse 使用）。
    與桌面版 _export_excel() 邏輯相同，不修改桌面版程式碼。
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    filters: dict = {}
    if date_from: filters["date_from"] = date_from
    if date_to:   filters["date_to"]   = date_to
    if type_ and type_ != "all": filters["type"] = type_

    app_title = db.get_setting("app_title", "個人記帳簿").strip() or "個人記帳簿"

    # ── Sheet 1: 明細 ─────────────────────────────────────────────────────────
    if include_detail:
        ws = wb.create_sheet("明細記錄")
        headers = ["日期", "收支說明", "收入", "支出", "餘額", "傳票編號"]

        ws.append([app_title] + [""] * (len(headers) - 1))
        end_col = get_column_letter(len(headers))
        ws.merge_cells(f"A1:{end_col}1")
        ws.cell(1, 1).font      = Font(bold=True, size=16)
        ws.cell(1, 1).alignment = _CENTER

        ws.append(["收支明細表"] + [""] * (len(headers) - 1))
        ws.merge_cells(f"A2:{end_col}2")
        ws.cell(2, 1).font      = Font(bold=True, size=14)
        ws.cell(2, 1).alignment = _CENTER

        range_txt = f"{date_from} ～ {date_to}" if date_from and date_to else "全部"
        ws.append([range_txt] + [""] * (len(headers) - 1))
        ws.merge_cells(f"A3:{end_col}3")
        ws.cell(3, 1).font      = Font(bold=True, size=12)
        ws.cell(3, 1).alignment = _CENTER

        ws.append([])
        ws.append(headers)
        _hdr(ws, ws.max_row, len(headers))

        rows = sorted(db.get_transactions(filters or None),
                      key=lambda r: (r["date"], r["id"]))
        prior = db.get_balance_before(
            filters.get("date_from", "2000-01-01"),
            filters.get("type"), filters.get("category")
        )
        running = prior

        if date_from:
            ws.append([date_from, "前期轉入",
                        prior if prior >= 0 else None,
                        -prior if prior < 0 else None,
                        running, ""])
            fr = ws.max_row
            fill = _INC_FILL if prior >= 0 else _EXP_FILL
            for c in range(1, len(headers) + 1):
                cell = ws.cell(fr, c)
                cell.fill = fill
                if c in (3, 4, 5): cell.number_format = "#,##0"
                cell.alignment = _CENTER
                cell.border = _BORDER

        for r in rows:
            is_inc = r["type"] == "income"
            desc   = (r.get("note") or r["category"]).strip()
            running += r["amount"] if is_inc else -r["amount"]
            ws.append([
                r["date"], desc,
                r["amount"] if is_inc else None,
                r["amount"] if not is_inc else None,
                running,
                r.get("voucher_no") or "",
            ])
            row_idx = ws.max_row
            fill = _INC_FILL if is_inc else _EXP_FILL
            font = _INC_FONT if is_inc else _EXP_FONT
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row_idx, c)
                cell.fill = fill
                if c in (3, 4, 5):
                    cell.font = font
                    cell.number_format = "#,##0"
                cell.alignment = _CENTER
                cell.border = _BORDER

        _autowidth(ws)

    # ── Sheet 2: 年度統計 ─────────────────────────────────────────────────────
    if include_year:
        ws = wb.create_sheet("年度統計")
        ws.append(["年度", "收入 (NT$)", "支出 (NT$)", "收支差 (NT$)", "筆數"])   # S10：收入不含期初、支出不含交接轉出
        _hdr(ws, 1, 5)
        year_rows = db.get_stats_by_year()
        for r in year_rows:
            inc = r.get("income", 0) or 0
            exp = r.get("expense", 0) or 0
            ws.append([r["year"], inc, exp, inc - exp, r.get("count", 0)])
        if year_rows:
            g_inc = sum((r.get("income", 0) or 0)  for r in year_rows)
            g_exp = sum((r.get("expense", 0) or 0) for r in year_rows)
            ws.append(["總計", g_inc, g_exp, g_inc - g_exp,
                        sum(r.get("count", 0) for r in year_rows)])
            tr = ws.max_row
            for c in range(1, 6):
                cell = ws.cell(tr, c)
                cell.fill = _TOTAL_FILL; cell.font = _TOTAL_FONT
                cell.alignment = _CENTER
                if c in (2, 3, 4): cell.number_format = "#,##0"
        _autowidth(ws)

    # ── Sheet 3: 月份統計 ─────────────────────────────────────────────────────
    if include_month:
        ws = wb.create_sheet("月份統計")
        ws.append(["月份", "收入 (NT$)", "支出 (NT$)", "收支差 (NT$)", "筆數"])   # S10
        _hdr(ws, 1, 5)
        for r in db.get_stats_by_month(year_filter or None):
            inc = r.get("income", 0) or 0
            exp = r.get("expense", 0) or 0
            ws.append([r["month"], inc, exp, inc - exp, r.get("count", 0)])
        _autowidth(ws)

    # ── Sheet 4: 分類統計 ─────────────────────────────────────────────────────
    if include_category:
        ws = wb.create_sheet("分類統計")
        ws.append(["類型", "分類", "金額 (NT$)", "佔比"])
        _hdr(ws, 1, 4)
        for type_label, type_key, fill_color in [
            ("支出", "expense", "FEF2F2"),
            ("收入", "income",  "ECFDF5"),
        ]:
            cats  = db.get_stats_by_category(type_key, filters or None)
            total = sum(v for _, v in cats)
            for cat, val in cats:
                pct = f"{val/total*100:.1f}%" if total else "0%"
                rn  = ws.max_row + 1
                ws.append([type_label, cat, val, pct])
                ws.cell(rn, 1).fill = PatternFill("solid", fgColor=fill_color)
        _autowidth(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_rent_xlsx(
    db: ReadOnlyDB,
    start_year: int, start_month: int,
    end_year:   int, end_month:   int,
) -> bytes:
    """
    產生月費繳交明細 Excel，格式與桌面版一致，但停用人員獨立於 Sheet 3：

    Sheet 1「月費繳交明細表」：啟用人員（V / X / 灰底斜線）
    Sheet 2「月費金額明細」  ：啟用人員（paid 顯示金額，淡綠底）
    Sheet 3「停用人員紀錄」  ：停用人員（範圍內有收費月份或有記錄者，相同格式）
    S2：不在收費期間（起算月份前／最後收費月份後）的格子為灰底。
    """
    import datetime as _dt
    from openpyxl.styles import Font as XLFont, PatternFill, Alignment, Border, Side

    GRADE_ORDER = ["一年級","二年級","三年級","四年級","五年級","六年級"]

    all_tenants = db.get_tenants(active_only=False)

    # ── 月份清單 ──────────────────────────────────────────────────────────────
    months: list[tuple[int, int]] = []
    y, m = start_year, start_month
    while (y, m) <= (end_year, end_month) and len(months) < 120:
        months.append((y, m))
        m += 1
        if m > 12:
            m = 1; y += 1

    range_start_ym = start_year * 100 + start_month
    range_end_ym   = end_year   * 100 + end_month

    # ── 範圍內是否有紀錄 ─────────────────────────────────────────────────────────
    def has_record_in_range(t):
        row = db._db.conn.execute(
            "SELECT COUNT(*) AS cnt FROM rent_payments "
            "WHERE tenant_id=? AND (year*100+month) BETWEEN ? AND ?",
            (t["id"], range_start_ym, range_end_ym)
        ).fetchone()
        return (row["cnt"] > 0) if row else False

    def grade_key_active(t):
        g = t.get("grade", "")
        grade_idx = GRADE_ORDER.index(g) if g in GRADE_ORDER else 99
        return (-grade_idx if g in GRADE_ORDER else 99, t["name"])

    # S2：範圍內有任一收費月份（起算月份～最後收費月份）或有紀錄者才列出
    _billable = db._db.is_tenant_billable

    def in_range(t):
        return any(_billable(t, y_, m_) for (y_, m_) in months) or has_record_in_range(t)

    active_tenants   = sorted([t for t in all_tenants
                                if t.get("active", 1) and in_range(t)],
                               key=grade_key_active)
    inactive_tenants = sorted([t for t in all_tenants
                                if not t.get("active", 1) and in_range(t)],
                               key=grade_key_active)

    # ── 預先批次讀取所有繳費記錄 ──────────────────────────────────────────────
    pay_rows = db._db.conn.execute(
        "SELECT * FROM rent_payments "
        "WHERE (year*100+month) BETWEEN ? AND ?",
        (range_start_ym, range_end_ym)
    ).fetchall()
    pay_map: dict[tuple, str] = {
        (r["tenant_id"], r["year"], r["month"]): r["status"] for r in pay_rows
    }
    # S3：金額明細使用「該月實際金額」（已繳／免繳為固定金額）
    pay_row_map = {(r["tenant_id"], r["year"], r["month"]): r for r in pay_rows}

    # ── 樣式定義 ──────────────────────────────────────────────────────────────
    title_font      = XLFont(name="新細明體", size=14, bold=True)
    header_font     = XLFont(name="新細明體", size=10, bold=True)
    cell_font       = XLFont(name="新細明體", size=10)
    inactive_font   = XLFont(name="新細明體", size=10, color="888888")
    paid_font       = XLFont(name="新細明體", size=10, color="065F46", bold=True)
    paid_font_inact = XLFont(name="新細明體", size=10, color="888888", bold=False)

    header_fill = PatternFill("solid", fgColor="D9D9D9")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    grey_fill   = PatternFill("solid", fgColor="CCCCCC")
    paid_fill   = PatternFill("solid", fgColor="ECFDF5")

    thin   = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def exempt_border():
        return Border(left=thin, right=thin, top=thin, bottom=thin,
                      diagonalUp=True, diagonal=thin)   # S4：免繳斜線「/」（原 diagonalDown 為「\」）

    center   = Alignment(horizontal="center", vertical="center")
    left_aln = Alignment(horizontal="left",   vertical="center")

    total_cols = 3 + len(months)

    # ══════════════════════════════════════════════════════════════════════════
    # 共用函式：將一組人員列表寫入一個工作表
    # ══════════════════════════════════════════════════════════════════════════

    def _write_sheet_vx(ws, title: str, header_label: str,
                        tenants: list, is_amount: bool = False):
        """
        寫入一張月費工作表（V/X 格式或金額格式）。
        is_amount=False → Sheet 1/3 格式（V/X）
        is_amount=True  → Sheet 2 格式（付款金額）
        """
        # Row 1：大標題
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        tc = ws.cell(1, 1, title)
        tc.font = title_font; tc.alignment = center; tc.border = border

        # Row 2：固定欄頭（編號/年級/姓名 跨 2 列）+ 月份群組頭
        for col, lbl in enumerate(["編號", "年級", "姓名"], start=1):
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
            c = ws.cell(2, col, lbl)
            c.font = header_font; c.fill = header_fill
            c.alignment = center; c.border = border
            ws.cell(3, col).border = border

        if months:
            ws.merge_cells(start_row=2, start_column=4,
                           end_row=2, end_column=total_cols)
        fh = ws.cell(2, 4, header_label)
        fh.font = header_font; fh.fill = header_fill
        fh.alignment = center; fh.border = border

        # Row 3：各月份（date 物件，number_format mm-yyyy）
        for j, (yr, mo) in enumerate(months):
            c = ws.cell(3, 4 + j, _dt.date(yr, mo, 1))
            c.number_format = "mm-yyyy"
            c.font = header_font; c.fill = header_fill
            c.alignment = center; c.border = border

        # 資料列
        idx = 0
        for i, t in enumerate(tenants):
            is_active    = bool(t.get("active", 1))
            use_font     = cell_font if is_active else inactive_font
            idx += 1
            row = 4 + i

            ws.cell(row, 1, idx).font = use_font
            ws.cell(row, 1).alignment = center
            ws.cell(row, 1).border    = border

            ws.cell(row, 2, t.get("grade", "")).font = use_font
            ws.cell(row, 2).alignment = center
            ws.cell(row, 2).border    = border

            ws.cell(row, 3, t["name"]).font = use_font
            ws.cell(row, 3).alignment = left_aln
            ws.cell(row, 3).border    = border

            for j, (yr, mo) in enumerate(months):
                col    = 4 + j
                status = pay_map.get((t["id"], yr, mo))
                billable = _billable(t, yr, mo)   # S2：不在收費期間＝灰底
                c = ws.cell(row, col)
                c.alignment = center

                if status == "paid":
                    if is_amount:
                        c.value         = db._db.fee_for_month(
                            t, yr, mo, pay_row_map.get((t["id"], yr, mo)))
                        c.number_format = "#,##0"
                        c.fill          = paid_fill
                        c.font          = paid_font if is_active else paid_font_inact
                    else:
                        c.value = "V"
                        c.fill  = white_fill
                        c.font  = use_font
                    c.border = border
                elif status == "exempt":
                    c.value  = ""
                    c.fill   = grey_fill
                    c.font   = use_font
                    c.border = exempt_border()
                elif status == "pending":
                    c.value  = "" if is_amount else "X"
                    c.fill   = white_fill if billable else grey_fill
                    c.font   = use_font
                    c.border = border
                else:  # 無記錄
                    c.value  = ""
                    c.fill   = white_fill if billable else grey_fill
                    c.font   = use_font
                    c.border = border

        # 欄寬
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 14
        col_w = 11 if is_amount else 9
        for j in range(len(months)):
            ws.column_dimensions[get_column_letter(4 + j)].width = col_w

        # 列高
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 18
        for i in range(len(tenants)):
            ws.row_dimensions[4 + i].height = 18

    # ── 建立 Workbook ─────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # Sheet 1：月費繳交明細表（啟用人員，V/X）
    ws1 = wb.active
    ws1.title = "月費繳交明細表"
    _write_sheet_vx(ws1, "月費繳交明細表", "繳費年月紀錄",
                    active_tenants, is_amount=False)

    # Sheet 2：月費金額明細（啟用人員，金額）
    ws2 = wb.create_sheet("月費金額明細")
    _write_sheet_vx(ws2, "月費金額明細表", "繳費年月金額紀錄",
                    active_tenants, is_amount=True)

    # Sheet 3：停用人員紀錄（有記錄者才建立）
    if inactive_tenants:
        ws3 = wb.create_sheet("停用人員紀錄")
        _write_sheet_vx(ws3, "停用人員紀錄", "繳費年月紀錄",
                        inactive_tenants, is_amount=False)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
